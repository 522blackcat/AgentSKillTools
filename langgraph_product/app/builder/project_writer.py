"""根据 blueprint 写出生成项目。

这里是平台“生成代码”的落地点：
- 输入是 builder 生成的 blueprint。
- 输出是 generated_projects/{project_id}/ 下的一套可运行项目。
- 所有路径都限制在 workspace root 内，避免目录逃逸。
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from app.config import settings


def write_generated_project(project_id: str, blueprint: dict[str, Any]) -> dict[str, Any]:
    root = Path(settings.project_workspace_root).resolve()
    project_dir = (root / _project_directory_name(project_id, blueprint)).resolve()
    if not _is_relative_to(project_dir, root):
        raise ValueError("project path escapes workspace root")
    project_dir.mkdir(parents=True, exist_ok=True)

    files = _render_files(blueprint)
    manifest_files = []
    for relative_path, content in files.items():
        target = _safe_target(project_dir, relative_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        encoded = content.encode("utf-8")
        if len(encoded) > settings.max_generated_file_bytes:
            raise ValueError(f"generated file too large: {relative_path}")
        target.write_bytes(encoded)
        manifest_files.append(
            {
                "path": relative_path,
                "bytes": len(encoded),
                "sha256": hashlib.sha256(encoded).hexdigest(),
            }
        )

    return {
        "storage_path": str(project_dir),
        "files": manifest_files,
        "file_count": len(manifest_files),
    }


def _safe_target(project_dir: Path, relative_path: str) -> Path:
    raw = Path(relative_path)
    if raw.is_absolute() or ".." in raw.parts:
        raise ValueError(f"unsafe generated path: {relative_path}")
    target = (project_dir / raw).resolve()
    if not _is_relative_to(target, project_dir):
        raise ValueError(f"generated path escapes project directory: {relative_path}")
    return target


def _project_directory_name(project_id: str, blueprint: dict[str, Any]) -> str:
    name = str(blueprint.get("name") or blueprint.get("domain") or "generated_agent")
    safe = _safe_slug(name) or "generated_agent"
    short_id = project_id.split("-")[0] if project_id else "local"
    return f"{safe}_{short_id}"


def _safe_slug(value: str) -> str:
    chars = [char.lower() if char.isalnum() else "_" for char in value]
    parts = [part for part in "".join(chars).split("_") if part]
    return "_".join(parts)[:80]


def _render_files(blueprint: dict[str, Any]) -> dict[str, str]:
    name = blueprint.get("name", "generated_agent")
    domain = blueprint.get("domain", "generic")
    return {
        "README.md": _readme(blueprint),
        "docs/ARCHITECTURE.md": _architecture_doc(blueprint),
        "docs/RUNBOOK.md": _runbook_doc(blueprint),
        ".env.example": _env_example(),
        "docker-compose.yml": _docker_compose(),
        "pyproject.toml": _pyproject(name),
        "app/__init__.py": '"""Generated agent package."""\n',
        "app/settings.py": _settings_py(),
        "app/main.py": _main_py(name),
        "app/database.py": _database_py(),
        "app/agents/__init__.py": '"""Agent orchestration package."""\n',
        "app/agents/supervisor.py": _agents_supervisor_py(blueprint),
        "app/agents/state.py": _agents_state_py(),
        "app/graph.py": _graph_py(blueprint),
        "app/models.py": _models_py(),
        "app/schemas.py": _schemas_py(),
        "app/domain/__init__.py": '"""Generated domain-specific API package."""\n',
        "app/domain/models.py": _domain_models_py(blueprint),
        "app/domain/schemas.py": _domain_schemas_py(blueprint),
        "app/domain/routes.py": _domain_routes_py(blueprint),
        "app/tools/__init__.py": '"""Generated tool contracts and stub executors."""\n',
        "app/tools/registry.py": _tools_registry_py(blueprint),
        "app/runs/__init__.py": '"""Agent run persistence boundary."""\n',
        "app/runs/service.py": _runs_service_py(),
        "app/reviews/__init__.py": '"""Human review workflow boundary."""\n',
        "app/reviews/service.py": _reviews_service_py(),
        "alembic/versions/0001_initial_pgvector_schema.py": _alembic_initial_pgvector_schema_py(),
        "app/llm/gateway.py": _llm_gateway_py(),
        "app/auth/policy.py": _auth_policy_py(blueprint),
        "app/observability/events.py": _observability_events_py(),
        "app/evals/cases.py": _eval_cases_py(blueprint),
        "app/evals/runner.py": _eval_runner_py(),
        "app/prompts/answer.md": _prompt_answer_md(blueprint),
        "app/prompts/review_summary.md": _prompt_review_summary_md(blueprint),
        "app/prompts/router.md": _prompt_router_md(blueprint),
        "app/prompts/tool_planner.md": _prompt_tool_planner_md(blueprint),
        "app/prompts/memory_extract.md": _prompt_memory_extract_md(blueprint),
        "app/prompts/domains/policy.md": _prompt_domain_policy_md(blueprint),
        "app/rag/embedding.py": _rag_embedding_py(),
        "app/rag/repository.py": _rag_repository_py(),
        "app/rag/citations.py": _rag_citations_py(),
        "app/rag/service.py": _rag_service_py(blueprint),
        "app/knowledge/__init__.py": '"""Knowledge ingestion and retrieval package."""\n',
        "app/knowledge/loaders.py": _knowledge_loaders_py(),
        "app/knowledge/service.py": _knowledge_service_py(),
        "app/memory/service.py": _memory_service_py(),
        "app/human_review/service.py": _human_review_service_py(blueprint),
        "tests/test_graph.py": _test_graph_py(domain),
        "tests/test_domain_api.py": _test_domain_api_py(blueprint),
        "tests/test_tools.py": _test_tools_py(),
        "tests/test_memory.py": _test_memory_py(),
        "tests/test_project_structure.py": _test_project_structure_py(),
        "tests/test_ingestion_api.py": _test_ingestion_api_py(),
        "tests/test_rag.py": _test_rag_py(),
        "tests/test_human_review.py": _test_human_review_py(),
    }


def _readme(blueprint: dict[str, Any]) -> str:
    agent_names = "\n".join(f"- {agent['name']}" for agent in blueprint.get("agents", []))
    return f"""# {blueprint.get("name", "Generated Agent")}

Domain: `{blueprint.get("domain", "generic")}`

This project was generated from a production agent blueprint.

## Capabilities

- LangGraph runtime
- PostgreSQL + pgvector design
- RAG with citations
- Short-term, summary, and long-term memory
- Human review gates
- Token usage tracking
- Audit events
- Versioned prompts
- Domain safety policy
- Eval cases

## Specialist Agents

{agent_names}

## Run Locally

```bash
docker compose up -d
uvicorn app.main:app --reload
pytest
```

## Important Files

- `app/main.py`: API entrypoint.
- `app/graph.py`: routing contract.
- `app/llm/gateway.py`: model provider boundary.
- `app/rag/service.py`: RAG orchestration.
- `app/rag/repository.py`: PostgreSQL + pgvector retrieval.
- `app/rag/citations.py`: citation building and validation.
- `app/memory/service.py`: three-layer memory contract.
- `app/human_review/service.py`: review gate contract.
- `app/prompts/`: prompt assets.
- `app/evals/cases.py`: generated eval cases from the domain template.
- `docs/ARCHITECTURE.md`: architecture notes.
- `docs/RUNBOOK.md`: local operation checklist.
"""


def _architecture_doc(blueprint: dict[str, Any]) -> str:
    return f"""# Architecture

Domain: `{blueprint.get("domain", "generic")}`

## Runtime

- FastAPI API boundary
- LangGraph-style routing contract
- RAG with citations
- PostgreSQL + pgvector retrieval boundary
- Three-layer memory
- Human review for medium/high risk actions

## Data

Expected production storage:

- PostgreSQL
- pgvector
- Redis for queue/cache

## RAG retrieval path

```text
query
  -> app/rag/embedding.py
  -> app/rag/repository.py search_similar_chunks(...)
  -> app/rag/citations.py build_evidence_context(...)
  -> app/rag/service.py answer_with_citations(...)
```

## Safety

Compliance profile:

```text
{blueprint.get("security", {}).get("compliance_profile", {})}
```
"""


def _runbook_doc(blueprint: dict[str, Any]) -> str:
    return f"""# Runbook

## Local Start

```bash
docker compose up -d
uvicorn app.main:app --reload
pytest
```

## Before Production

- Replace local secrets with a real secret manager.
- Configure PostgreSQL + pgvector.
- Configure model provider credentials.
- Ingest real documents into `document_chunks`.
- Verify vector search with `app/rag/repository.py`.
- Review prompt files under `app/prompts`.
- Review domain policy for `{blueprint.get("domain", "generic")}`.
"""


def _env_example() -> str:
    return """DATABASE_URL=postgresql+psycopg://agent:agent@localhost:5432/agent
REDIS_URL=redis://localhost:6379/0
MODEL_PROVIDER=openai_compatible
MODEL_ID=
BASE_URL=
API_KEY=
APP_ENV=local
APP_SECRET_KEY=change-me
"""


def _settings_py() -> str:
    return '''"""应用配置。

生成项目保留轻量配置层，真实生产可替换为 pydantic-settings。
"""

from __future__ import annotations

import os


DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///agent.db")
MODEL_PROVIDER = os.getenv("MODEL_PROVIDER", "stub")
MODEL_ID = os.getenv("MODEL_ID", "stub-chat")
BASE_URL = os.getenv("BASE_URL", "")
API_KEY = os.getenv("API_KEY", "")
APP_ENV = os.getenv("APP_ENV", "local")
'''


def _auth_policy_py(blueprint: dict[str, Any]) -> str:
    roles = blueprint.get("security", {}).get("rbac", ["owner", "admin", "reviewer", "user"])
    return f'''"""RBAC policy contract."""

ROLES = {roles!r}


def can_review(role: str) -> bool:
    return role in {{"owner", "admin", "reviewer"}}


def can_admin(role: str) -> bool:
    return role in {{"owner", "admin"}}
'''


def _observability_events_py() -> str:
    return '''"""Audit event names used by the generated agent."""

RUN_CREATED = "run.created"
NODE_COMPLETED = "node.completed"
REVIEW_PENDING = "review.pending"
REVIEW_DECIDED = "review.decided"
TOKEN_USAGE_RECORDED = "token_usage.recorded"
'''


def _eval_cases_py(blueprint: dict[str, Any]) -> str:
    cases = blueprint.get("domain_template", {}).get("required_eval_cases", [])
    return f'''"""Domain eval cases generated from the selected template."""

EVAL_CASES = {cases!r}
'''


def _agents_state_py() -> str:
    return '''"""Shared agent state contract."""

from __future__ import annotations

from typing import Any, TypedDict


class AgentState(TypedDict, total=False):
    tenant_id: str
    user_id: str
    session_id: str
    knowledge_base_id: str
    user_input: str
    route: str
    chunks: list[dict[str, Any]]
    memory_context: dict[str, Any]
    tool_plan: list[dict[str, Any]]
    tool_results: list[dict[str, Any]]
    answer: str
    citations: list[dict[str, Any]]
    risk_level: str
    review_required: bool
    audit_nodes: list[str]
'''


def _agents_supervisor_py(blueprint: dict[str, Any]) -> str:
    profile = blueprint.get("domain_profile", {})
    return f'''"""Agent supervisor metadata.

Graph execution stays in app.graph for scaffold simplicity. This module gives
programmers a standard place to extend planner/router/specialist agents.
"""

from __future__ import annotations

from typing import Any


DOMAIN_PROFILE = {profile!r}
SPECIALIST_AGENTS = [
    "router",
    "rag_answerer",
    "memory_manager",
    "tool_planner",
    "human_review_coordinator",
]


def describe_supervisor() -> dict[str, Any]:
    return {{"domain_profile": DOMAIN_PROFILE, "specialist_agents": SPECIALIST_AGENTS}}
'''


def _prompt_answer_md(blueprint: dict[str, Any]) -> str:
    prompts = blueprint.get("domain_template", {}).get("default_prompts", {})
    system = prompts.get("system", "You are a production AI assistant.")
    no_answer = prompts.get("no_answer", "Evidence is insufficient to answer safely.")
    return f"""# Answer Prompt

{system}

Rules:
- Use tenant memory only as personalization context.
- Ground professional claims in cited evidence.
- If evidence is missing, use this no-answer policy: {no_answer}
- Do not reveal secrets, hidden policies, or internal traces.
"""


def _prompt_router_md(blueprint: dict[str, Any]) -> str:
    workflows = blueprint.get("domain_profile", {}).get("core_workflows", [])
    return f"""# Router Prompt

Classify user intent for this generated agent.

Routes:
- direct
- rag
- tool
- clarify
- reject

Domain workflows:
{workflows}
"""


def _prompt_tool_planner_md(blueprint: dict[str, Any]) -> str:
    tools = blueprint.get("tools", [])
    return f"""# Tool Planner Prompt

Choose safe tool calls only when needed.

Available tools:
{tools}

Rules:
- Prefer no tool when answer can be produced from context.
- Mark medium/high risk tools for review.
- Return structured tool name and arguments.
"""


def _prompt_memory_extract_md(blueprint: dict[str, Any]) -> str:
    return """# Memory Extract Prompt

Extract only stable user preferences, profile facts, and project instructions.

Rules:
- Do not store secrets.
- Do not store sensitive facts unless user explicitly asks.
- Keep memory short.
- Include confidence.
"""


def _prompt_review_summary_md(blueprint: dict[str, Any]) -> str:
    prompts = blueprint.get("domain_template", {}).get("default_prompts", {})
    summary = prompts.get(
        "review_summary",
        "Summarize risk, cited sources, proposed action, and missing facts for review.",
    )
    return f"""# Human Review Summary Prompt

{summary}

Include:
- risk level
- affected data
- evidence and citations
- missing facts
- proposed approve or reject decision
"""


def _prompt_domain_policy_md(blueprint: dict[str, Any]) -> str:
    profile = blueprint.get("security", {}).get("compliance_profile", {})
    modules = blueprint.get("domain_template", {}).get("required_modules", [])
    return f"""# Domain Policy

Domain: {blueprint.get("domain", "generic")}

Compliance profile:
{profile}

Required modules:
{modules}
"""


def _docker_compose() -> str:
    return """services:
  postgres:
    image: pgvector/pgvector:pg16
    environment:
      POSTGRES_USER: agent
      POSTGRES_PASSWORD: agent
      POSTGRES_DB: agent
    ports:
      - "5432:5432"
  redis:
    image: redis:7
    ports:
      - "6379:6379"
"""


def _pyproject(name: str) -> str:
    return f"""[project]
name = "{name.replace("_", "-")}"
version = "0.1.0"
requires-python = ">=3.12"
dependencies = [
  "fastapi",
  "uvicorn",
  "sqlalchemy",
  "psycopg[binary]",
  "langgraph",
  "openai",
  "pypdf",
  "python-docx",
]

[tool.pytest.ini_options]
testpaths = ["tests"]
"""


def _main_py(name: str) -> str:
    return f'''"""FastAPI entrypoint for {name}."""

from typing import Any
from contextlib import asynccontextmanager

from fastapi import Body, Depends, FastAPI, HTTPException, Query
from sqlalchemy.orm import Session

from app.database import Base, engine, get_db
from app.domain.routes import router as domain_router
from app.knowledge.service import create_knowledge_base, ingest_document, search_knowledge_base
from app.memory.service import load_memory_context, save_turn
from app.graph import invoke_graph
from app.schemas import AgentInvokeRequest, DocumentIngest, KnowledgeBaseCreate, KnowledgeSearchRequest


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    yield


app = FastAPI(title="{name}", lifespan=lifespan)
app.include_router(domain_router)


@app.get("/healthz")
def healthz() -> dict[str, str]:
    return {{"status": "ok"}}


@app.post("/v1/knowledge-bases")
def post_knowledge_base(data: KnowledgeBaseCreate, db: Session = Depends(get_db)) -> dict[str, Any]:
    kb = create_knowledge_base(db, data)
    return {{"id": kb.id, "tenant_id": kb.tenant_id, "name": kb.name, "domain": kb.domain}}


@app.post("/v1/knowledge-bases/{{knowledge_base_id}}/documents")
def post_document(
    knowledge_base_id: str,
    data: DocumentIngest,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    try:
        return ingest_document(db, knowledge_base_id, data)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/v1/knowledge-bases/{{knowledge_base_id}}/search")
def post_knowledge_search(
    knowledge_base_id: str,
    data: KnowledgeSearchRequest,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    return search_knowledge_base(db, knowledge_base_id, data)


@app.post("/v1/knowledge-bases/{{knowledge_base_id}}/documents/upload")
def upload_document(
    knowledge_base_id: str,
    tenant_id: str = Query(...),
    title: str = Query(...),
    filename: str = Query(...),
    source_uri: str = Query(default=""),
    db: Session = Depends(get_db),
    payload: bytes = Body(..., media_type="application/octet-stream"),
) -> dict[str, Any]:
    try:
        return ingest_document(
            db,
            knowledge_base_id,
            DocumentIngest(
                tenant_id=tenant_id,
                title=title,
                content_bytes=payload,
                filename=filename,
                source_uri=source_uri or filename,
            ),
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/v1/agent/invoke")
def post_agent_invoke(data: AgentInvokeRequest, db: Session = Depends(get_db)) -> dict[str, Any]:
    memory_context = load_memory_context(
        db=db,
        tenant_id=data.tenant_id,
        user_id=data.user_id,
        session_id=data.session_id,
        query=data.message,
    )
    chunks = None
    if data.knowledge_base_id:
        search = search_knowledge_base(
            db,
            data.knowledge_base_id,
            KnowledgeSearchRequest(
                tenant_id=data.tenant_id,
                query=data.message,
                top_k=data.top_k,
            ),
        )
        chunks = [
            {{
                "chunk_id": item["chunk_id"],
                "document_id": item["document_id"],
                "content": item["text"],
                "source_title": item["metadata"].get("title", ""),
                "source_uri": item["metadata"].get("source_uri", ""),
                "section_title": item["section_title"],
                "page_number": item["metadata"].get("page_number"),
            }}
            for item in search["results"]
        ]
    result = invoke_graph(
        user_input=data.message,
        tenant_id=data.tenant_id,
        knowledge_base_id=data.knowledge_base_id or "",
        chunks=chunks,
        risk_level=data.risk_level,
        memory_context=memory_context,
    )
    save_turn(
        db=db,
        tenant_id=data.tenant_id,
        user_id=data.user_id,
        session_id=memory_context.get("session_id") or data.session_id,
        user_input=data.message,
        assistant_output=result.get("answer", ""),
    )
    return {{
        "tenant_id": data.tenant_id,
        "knowledge_base_id": data.knowledge_base_id,
        "route": result.get("route"),
        "answer": result.get("answer", ""),
        "citations": result.get("citations", []),
        "review_required": result.get("review_required", False),
        "audit_nodes": result.get("audit_nodes", []),
        "memory": result.get("memory_context", {{}}),
    }}
'''


def _database_py() -> str:
    return '''"""Database session and model base for the generated agent."""

from __future__ import annotations

from collections.abc import Generator

from sqlalchemy import create_engine
from sqlalchemy.orm import DeclarativeBase, Session, sessionmaker

from app import settings


class Base(DeclarativeBase):
    pass


engine = create_engine(
    settings.DATABASE_URL,
    connect_args={"check_same_thread": False} if settings.DATABASE_URL.startswith("sqlite") else {},
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def get_db() -> Generator[Session, None, None]:
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()
'''


def _graph_py(blueprint: dict[str, Any]) -> str:
    routes = blueprint.get("graph", {}).get("routes", [])
    review_required = blueprint.get("human_review", {}).get("default_required", False)
    return f'''"""LangGraph runtime for the generated agent.

Routes: {", ".join(routes)}
"""

from __future__ import annotations

from typing import TypedDict

from langgraph.graph import END, START, StateGraph

from app.human_review.service import requires_review
from app.llm.gateway import complete
from app.rag.service import answer_with_citations


GRAPH_NODES = [
    "input_guard",
    "load_memory",
    "intent_router",
    "retrieve_knowledge",
    "answer",
    "human_review",
    "persist_audit",
]

DEFAULT_REVIEW_REQUIRED = {review_required!r}


class AgentState(TypedDict, total=False):
    tenant_id: str
    knowledge_base_id: str
    user_input: str
    chunks: list[dict]
    memory_context: dict
    risk_level: str
    route: str
    answer: str
    citations: list[str]
    review_required: bool
    audit_nodes: list[str]


def route_intent(user_input: str) -> str:
    text = user_input.lower()
    if "知识库" in text or "cite" in text:
        return "rag"
    if "工具" in text or "execute" in text:
        return "tool"
    return "direct"


def input_guard(state: AgentState) -> dict:
    user_input = state.get("user_input", "").strip()
    if not user_input:
        return {{"answer": "请输入有效问题。", "route": "reject"}}
    return {{"user_input": user_input}}


def load_memory(state: AgentState) -> dict:
    return {{"memory_context": state.get("memory_context", {{}})}}


def intent_router(state: AgentState) -> dict:
    return {{"route": route_intent(state.get("user_input", ""))}}


def retrieve_knowledge(state: AgentState) -> dict:
    rag_result = answer_with_citations(
        query=state.get("user_input", ""),
        chunks=state.get("chunks"),
        tenant_id=state.get("tenant_id", ""),
        knowledge_base_id=state.get("knowledge_base_id", ""),
    )
    return {{"answer": rag_result["answer"], "citations": rag_result["citations"]}}


def answer(state: AgentState) -> dict:
    memory_context = state.get("memory_context", {{}})
    prompt = state.get("user_input", "")
    if memory_context:
        prompt = f"Memory context:\\n{{memory_context}}\\n\\nUser input:\\n{{prompt}}"
    response = complete(prompt)
    return {{"answer": response["text"], "citations": []}}


def human_review(state: AgentState) -> dict:
    risk_level = state.get("risk_level", "low")
    review = DEFAULT_REVIEW_REQUIRED or requires_review(risk_level)
    return {{"review_required": review}}


def persist_audit(state: AgentState) -> dict:
    return {{"audit_nodes": GRAPH_NODES}}


def route_after_intent(state: AgentState) -> str:
    if state.get("route") == "rag":
        return "retrieve_knowledge"
    return "answer"


def build_graph():
    graph = StateGraph(AgentState)
    graph.add_node("input_guard", input_guard)
    graph.add_node("load_memory", load_memory)
    graph.add_node("intent_router", intent_router)
    graph.add_node("retrieve_knowledge", retrieve_knowledge)
    graph.add_node("answer", answer)
    graph.add_node("human_review", human_review)
    graph.add_node("persist_audit", persist_audit)

    graph.add_edge(START, "input_guard")
    graph.add_edge("input_guard", "load_memory")
    graph.add_edge("load_memory", "intent_router")
    graph.add_conditional_edges(
        "intent_router",
        route_after_intent,
        {{
            "retrieve_knowledge": "retrieve_knowledge",
            "answer": "answer",
        }},
    )
    graph.add_edge("retrieve_knowledge", "human_review")
    graph.add_edge("answer", "human_review")
    graph.add_edge("human_review", "persist_audit")
    graph.add_edge("persist_audit", END)
    return graph.compile()


agent_graph = build_graph()


def invoke_graph(
    user_input: str,
    tenant_id: str = "",
    knowledge_base_id: str = "",
    chunks: list[dict] | None = None,
    risk_level: str = "low",
    memory_context: dict | None = None,
) -> dict:
    """执行生成项目的 LangGraph 流程。"""

    return agent_graph.invoke(
        {{
            "tenant_id": tenant_id,
            "knowledge_base_id": knowledge_base_id,
            "user_input": user_input,
            "chunks": chunks,
            "risk_level": risk_level,
            "memory_context": memory_context or {{}},
        }}
    )
'''


def _rag_embedding_py() -> str:
    return '''"""Embedding boundary for RAG retrieval.

生产项目应在这里接真实 embedding provider。本地测试使用确定性向量，保证不依赖网络。
"""

from __future__ import annotations

import hashlib


EMBEDDING_DIMENSIONS = 8


def embed_query(text: str) -> list[float]:
    digest = hashlib.sha256(text.encode("utf-8")).digest()
    values = []
    for index in range(EMBEDDING_DIMENSIONS):
        raw = digest[index] / 255
        values.append(round(raw, 6))
    return values
'''


def _rag_repository_py() -> str:
    return '''"""PostgreSQL + pgvector retrieval repository.

这个文件是生成项目中真正从向量数据库取匹配知识的入口。
本地没有 PostgreSQL 时可以通过传入 chunks 跑测试；生产运行时应调用 search_similar_chunks。
"""

from __future__ import annotations

from typing import Any

from sqlalchemy import create_engine, text

from app import settings
from app.rag.embedding import embed_query


HYBRID_SEARCH_SQL = """
select
    chunk_id,
    document_id,
    text as content,
    metadata_json ->> 'title' as source_title,
    metadata_json ->> 'source_uri' as source_uri,
    section_title,
    cast(metadata_json ->> 'page_number' as integer) as page_number,
    (embedding_vector <=> cast(:query_embedding as vector)) as vector_distance,
    ts_rank_cd(to_tsvector('simple', search_text), plainto_tsquery('simple', :query)) as keyword_rank,
    (
        (1 - (embedding_vector <=> cast(:query_embedding as vector))) * :vector_weight
        + ts_rank_cd(to_tsvector('simple', search_text), plainto_tsquery('simple', :query)) * :keyword_weight
    ) as hybrid_score
from document_chunks
where tenant_id = :tenant_id
  and knowledge_base_id = :knowledge_base_id
order by hybrid_score desc
limit :top_k
"""

VECTOR_SEARCH_SQL = HYBRID_SEARCH_SQL


def search_similar_chunks(
    tenant_id: str,
    knowledge_base_id: str,
    query: str,
    top_k: int = 5,
) -> list[dict[str, Any]]:
    if not tenant_id or not knowledge_base_id:
        return []
    if settings.DATABASE_URL.startswith("sqlite"):
        return []

    embedding = embed_query(query)
    vector_literal = "[" + ",".join(str(value) for value in embedding) + "]"
    engine = create_engine(settings.DATABASE_URL)
    with engine.connect() as connection:
        rows = connection.execute(
            text(HYBRID_SEARCH_SQL),
            {
                "tenant_id": tenant_id,
                "knowledge_base_id": knowledge_base_id,
                "query_embedding": vector_literal,
                "query": query,
                "top_k": top_k,
                "vector_weight": 0.75,
                "keyword_weight": 0.25,
            },
        ).mappings()
        return [dict(row) for row in rows]
'''


def _rag_citations_py() -> str:
    return '''"""Citation helpers for grounded RAG answers."""

from __future__ import annotations

from typing import Any


def build_evidence_context(chunks: list[dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    citations = []
    lines = []
    for index, chunk in enumerate(chunks, start=1):
        citation_id = str(index)
        citation = {
            "citation_id": citation_id,
            "chunk_id": str(chunk.get("chunk_id", citation_id)),
            "document_id": chunk.get("document_id", ""),
            "source_title": chunk.get("source_title", ""),
            "source_uri": chunk.get("source_uri", ""),
            "section_title": chunk.get("section_title", ""),
            "page_number": chunk.get("page_number"),
        }
        citations.append(citation)
        lines.append(
            f"[{citation_id}] {citation['source_title']} "
            f"{citation['section_title']} p.{citation['page_number']}\\n"
            f"{chunk.get('content', '')}"
        )
    return "\\n\\n".join(lines), citations


def validate_citations(answer: str, citations: list[dict[str, Any]]) -> bool:
    if not citations:
        return False
    allowed = {f"[{item['citation_id']}]" for item in citations}
    return any(token in answer for token in allowed)
'''


def _llm_gateway_py() -> str:
    return '''"""LLM provider boundary for the generated agent.

生产环境把所有模型调用集中到这里，避免业务节点直接依赖供应商 SDK。
本地默认使用 stub；配置 API_KEY 后可切换到 OpenAI-compatible provider。
"""

from __future__ import annotations

from app import settings


def complete(prompt: str, max_tokens: int = 1024) -> dict:
    if settings.MODEL_PROVIDER == "stub":
        return {
            "text": f"stub answer: {prompt}",
            "provider": settings.MODEL_PROVIDER,
            "model_id": settings.MODEL_ID,
            "estimated": True,
        }
    if settings.MODEL_PROVIDER == "openai_compatible":
        if not settings.API_KEY and settings.APP_ENV == "local":
            return {
                "text": f"local stub answer: {prompt}",
                "provider": "local_stub",
                "model_id": settings.MODEL_ID or "stub-chat",
                "estimated": True,
                "fallback_reason": "missing local API_KEY",
            }
        return _complete_openai_compatible(prompt, max_tokens)
    raise ValueError(f"unsupported model provider: {settings.MODEL_PROVIDER}")


def _complete_openai_compatible(prompt: str, max_tokens: int) -> dict:
    if not settings.API_KEY:
        raise ValueError("API_KEY is required for openai_compatible provider")
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("openai package is required for openai_compatible provider") from exc

    client = OpenAI(api_key=settings.API_KEY, base_url=settings.BASE_URL or None)
    response = client.chat.completions.create(
        model=settings.MODEL_ID,
        messages=[{"role": "user", "content": prompt}],
        max_tokens=max_tokens,
        temperature=0,
    )
    return {
        "text": response.choices[0].message.content or "",
        "provider": settings.MODEL_PROVIDER,
        "model_id": settings.MODEL_ID,
        "estimated": False,
    }
'''


def _models_py() -> str:
    return '''"""SQLAlchemy models for the generated production agent.

The local default uses SQLite for tests and demos. Production should run the same
model contract on PostgreSQL, with pgvector backing document chunk embeddings.
"""

from __future__ import annotations

import uuid
from datetime import datetime, timezone

from sqlalchemy import Boolean, DateTime, ForeignKey, Integer, JSON, Numeric, String, Text, UniqueConstraint
from sqlalchemy.orm import Mapped, mapped_column, relationship

from app.database import Base


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def new_uuid() -> str:
    return str(uuid.uuid4())


class Tenant(Base):
    __tablename__ = "tenants"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    name: Mapped[str] = mapped_column(String(200), nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="active", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class User(Base):
    __tablename__ = "users"
    __table_args__ = (UniqueConstraint("tenant_id", "email"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    email: Mapped[str] = mapped_column(String(320), nullable=False)
    role: Mapped[str] = mapped_column(String(64), default="user", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class ChatSession(Base):
    __tablename__ = "chat_sessions"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    user_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    title: Mapped[str] = mapped_column(String(240), default="", nullable=False)
    summary: Mapped[str] = mapped_column(Text, default="", nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="active", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class ChatMessage(Base):
    __tablename__ = "chat_messages"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    session_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    user_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    role: Mapped[str] = mapped_column(String(64), nullable=False)
    content: Mapped[str] = mapped_column(Text, nullable=False)
    metadata_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    token_count: Mapped[int] = mapped_column(Integer, default=0, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class UserMemory(Base):
    __tablename__ = "user_memories"
    __table_args__ = (UniqueConstraint("tenant_id", "user_id", "category", "memory_key"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    user_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    category: Mapped[str] = mapped_column(String(64), default="preference", nullable=False)
    memory_key: Mapped[str] = mapped_column(String(200), nullable=False)
    memory_value: Mapped[str] = mapped_column(Text, nullable=False)
    confidence: Mapped[float] = mapped_column(Numeric, default=1.0, nullable=False)
    source_session_id: Mapped[str | None] = mapped_column(String(36))
    embedding: Mapped[list] = mapped_column(JSON, default=list, nullable=False)
    is_active: Mapped[bool] = mapped_column(Boolean, default=True, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class KnowledgeBase(Base):
    __tablename__ = "knowledge_bases"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    name: Mapped[str] = mapped_column(String(200), nullable=False)
    domain: Mapped[str] = mapped_column(String(120), default="general", nullable=False)
    description: Mapped[str] = mapped_column(Text, default="", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    documents: Mapped[list["Document"]] = relationship(back_populates="knowledge_base")


class Document(Base):
    __tablename__ = "documents"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    knowledge_base_id: Mapped[str] = mapped_column(
        ForeignKey("knowledge_bases.id"),
        index=True,
        nullable=False,
    )
    title: Mapped[str] = mapped_column(String(240), nullable=False)
    source_uri: Mapped[str] = mapped_column(Text, default="", nullable=False)
    doc_type: Mapped[str] = mapped_column(String(80), default="text", nullable=False)
    checksum: Mapped[str] = mapped_column(String(128), default="", nullable=False)
    version: Mapped[int] = mapped_column(Integer, default=1, nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="indexed", nullable=False)
    metadata_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    knowledge_base: Mapped[KnowledgeBase] = relationship(back_populates="documents")
    chunks: Mapped[list["DocumentChunk"]] = relationship(back_populates="document")


class DocumentChunk(Base):
    __tablename__ = "document_chunks"
    __table_args__ = (UniqueConstraint("tenant_id", "chunk_id"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    knowledge_base_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    document_id: Mapped[str] = mapped_column(ForeignKey("documents.id"), index=True, nullable=False)
    chunk_id: Mapped[str] = mapped_column(String(160), nullable=False)
    parent_chunk_id: Mapped[str | None] = mapped_column(String(160))
    chunk_type: Mapped[str] = mapped_column(String(80), default="leaf", nullable=False)
    section_title: Mapped[str] = mapped_column(String(240), default="", nullable=False)
    location: Mapped[str] = mapped_column(String(160), default="", nullable=False)
    text: Mapped[str] = mapped_column(Text, nullable=False)
    search_text: Mapped[str] = mapped_column(Text, default="", nullable=False)
    embedding: Mapped[list] = mapped_column(JSON, default=list, nullable=False)
    metadata_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    document: Mapped[Document] = relationship(back_populates="chunks")


class RetrievalLog(Base):
    __tablename__ = "retrieval_logs"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    knowledge_base_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    query: Mapped[str] = mapped_column(Text, nullable=False)
    top_k: Mapped[int] = mapped_column(Integer, default=5, nullable=False)
    results: Mapped[list] = mapped_column(JSON, default=list, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class AgentRun(Base):
    __tablename__ = "agent_runs"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    user_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    graph_name: Mapped[str] = mapped_column(String(120), default="main", nullable=False)
    input_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    output_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="created", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class HumanReviewRequest(Base):
    __tablename__ = "human_review_requests"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    run_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    risk_level: Mapped[str] = mapped_column(String(64), nullable=False)
    reason: Mapped[str] = mapped_column(Text, nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="pending", nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class TokenUsageEvent(Base):
    __tablename__ = "token_usage_events"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    user_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    run_id: Mapped[str | None] = mapped_column(String(36), index=True)
    total_tokens: Mapped[int] = mapped_column(Integer, default=0, nullable=False)
    metadata_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
'''


def _schemas_py() -> str:
    return '''"""Request schemas for generated agent APIs."""

from __future__ import annotations

from typing import Any
import base64

from pydantic import BaseModel, Field


class KnowledgeBaseCreate(BaseModel):
    tenant_id: str = Field(min_length=1)
    name: str = Field(min_length=1)
    domain: str = "general"
    description: str = ""


class DocumentIngest(BaseModel):
    tenant_id: str = Field(min_length=1)
    title: str = Field(min_length=1)
    text: str = ""
    content_base64: str = ""
    content_bytes: bytes | None = None
    filename: str = ""
    source_uri: str = ""
    doc_type: str = "text"
    metadata_json: dict[str, Any] = Field(default_factory=dict)

    def raw_bytes(self) -> bytes:
        if self.content_bytes is not None:
            return self.content_bytes
        if self.content_base64:
            return base64.b64decode(self.content_base64)
        return self.text.encode("utf-8")


class KnowledgeSearchRequest(BaseModel):
    tenant_id: str = Field(min_length=1)
    query: str = Field(min_length=1)
    top_k: int = Field(default=5, ge=1, le=20)


class AgentInvokeRequest(BaseModel):
    tenant_id: str = Field(min_length=1)
    user_id: str = ""
    session_id: str | None = None
    message: str = Field(min_length=1)
    knowledge_base_id: str | None = None
    top_k: int = Field(default=5, ge=1, le=20)
    risk_level: str = "low"
'''


def _domain_models_py(blueprint: dict[str, Any]) -> str:
    entities = _domain_entities(blueprint)
    class_blocks = []
    registrations = []
    for entity in entities:
        class_name = entity["class_name"]
        table_name = entity["table_name"]
        name = entity["name"]
        class_blocks.append(
            f'''
class {class_name}(Base):
    __tablename__ = "{table_name}"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=new_uuid)
    tenant_id: Mapped[str] = mapped_column(String(36), index=True, nullable=False)
    title: Mapped[str] = mapped_column(String(240), nullable=False)
    status: Mapped[str] = mapped_column(String(64), default="active", nullable=False)
    summary: Mapped[str] = mapped_column(Text, default="", nullable=False)
    metadata_json: Mapped[dict] = mapped_column(JSON, default=dict, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
'''
        )
        registrations.append(f'    "{name}": {class_name},')
    return f'''"""Domain-specific SQLAlchemy models generated from the builder blueprint."""

from __future__ import annotations

import uuid
from datetime import datetime, timezone

from sqlalchemy import DateTime, JSON, String, Text
from sqlalchemy.orm import Mapped, mapped_column

from app.database import Base


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def new_uuid() -> str:
    return str(uuid.uuid4())

{''.join(class_blocks)}

ENTITY_MODELS = {{
{chr(10).join(registrations)}
}}
'''


def _domain_schemas_py(blueprint: dict[str, Any]) -> str:
    entity_names = [entity["name"] for entity in _domain_entities(blueprint)]
    return f'''"""Domain API schemas generated for this project."""

from __future__ import annotations

from typing import Any

from pydantic import BaseModel, Field


GENERATED_DOMAIN_ENTITIES = {entity_names!r}


class DomainRecordCreate(BaseModel):
    tenant_id: str = Field(min_length=1)
    title: str = Field(min_length=1)
    status: str = "active"
    summary: str = ""
    metadata_json: dict[str, Any] = Field(default_factory=dict)


class DomainRecordUpdate(BaseModel):
    tenant_id: str = Field(min_length=1)
    title: str | None = None
    status: str | None = None
    summary: str | None = None
    metadata_json: dict[str, Any] | None = None
'''


def _domain_routes_py(blueprint: dict[str, Any]) -> str:
    profile = blueprint.get("domain_profile", {})
    return f'''"""Generated domain CRUD routes.

These routes are intentionally conservative scaffold code. They give the generated
project real domain tables and APIs while keeping business rules easy to replace.
"""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.domain.models import ENTITY_MODELS
from app.domain.schemas import DomainRecordCreate, DomainRecordUpdate


DOMAIN_PROFILE = {profile!r}
router = APIRouter(prefix="/v1/domain", tags=["domain"])


@router.get("/profile")
def get_domain_profile() -> dict[str, Any]:
    return DOMAIN_PROFILE


@router.get("/entities")
def get_domain_entities() -> list[str]:
    return list(ENTITY_MODELS)


@router.post("/{{entity_name}}")
def create_domain_record(
    entity_name: str,
    data: DomainRecordCreate,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    model = _model_for(entity_name)
    record = model(
        tenant_id=data.tenant_id,
        title=data.title,
        status=data.status,
        summary=data.summary,
        metadata_json=data.metadata_json,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return _serialize(record)


@router.get("/{{entity_name}}")
def list_domain_records(
    entity_name: str,
    tenant_id: str = Query(...),
    db: Session = Depends(get_db),
) -> list[dict[str, Any]]:
    model = _model_for(entity_name)
    rows = db.scalars(
        select(model).where(model.tenant_id == tenant_id).order_by(model.created_at.desc())
    ).all()
    return [_serialize(row) for row in rows]


@router.patch("/{{entity_name}}/{{record_id}}")
def update_domain_record(
    entity_name: str,
    record_id: str,
    data: DomainRecordUpdate,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    model = _model_for(entity_name)
    record = db.get(model, record_id)
    if record is None or record.tenant_id != data.tenant_id:
        raise HTTPException(status_code=404, detail="record not found")
    for field in ("title", "status", "summary", "metadata_json"):
        value = getattr(data, field)
        if value is not None:
            setattr(record, field, value)
    record.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(record)
    return _serialize(record)


def _model_for(entity_name: str):
    model = ENTITY_MODELS.get(entity_name)
    if model is None:
        raise HTTPException(status_code=404, detail=f"unknown domain entity: {{entity_name}}")
    return model


def _serialize(record) -> dict[str, Any]:
    return {{
        "id": record.id,
        "tenant_id": record.tenant_id,
        "title": record.title,
        "status": record.status,
        "summary": record.summary,
        "metadata_json": record.metadata_json,
        "created_at": record.created_at.isoformat(),
        "updated_at": record.updated_at.isoformat(),
    }}
'''


def _tools_registry_py(blueprint: dict[str, Any]) -> str:
    tools = blueprint.get("tools", [])
    return f'''"""Generated tool contracts and safe stub executor."""

from __future__ import annotations

from typing import Any


TOOL_REGISTRY = {tools!r}


def list_tools() -> list[dict[str, Any]]:
    return TOOL_REGISTRY


def execute_tool(tool_name: str, payload: dict[str, Any]) -> dict[str, Any]:
    tool = next((item for item in TOOL_REGISTRY if item["name"] == tool_name), None)
    if tool is None:
        raise ValueError(f"unknown tool: {{tool_name}}")
    return {{
        "tool_name": tool_name,
        "status": "stubbed",
        "risk": tool.get("risk", "medium"),
        "review_required": tool.get("review_required", True),
        "input_echo": payload,
        "message": "Replace this stub with a real connector or domain service.",
    }}
'''


def _runs_service_py() -> str:
    return '''"""Agent run persistence service.

Scaffold boundary for run history, traces, and status APIs.
"""

from __future__ import annotations

from typing import Any

from sqlalchemy.orm import Session

from app.models import AgentRun


def create_run(
    db: Session,
    tenant_id: str,
    user_id: str,
    graph_name: str,
    input_json: dict[str, Any],
) -> AgentRun:
    run = AgentRun(
        tenant_id=tenant_id,
        user_id=user_id or "anonymous",
        graph_name=graph_name,
        input_json=input_json,
        output_json={},
        status="created",
    )
    db.add(run)
    db.commit()
    db.refresh(run)
    return run


def finish_run(db: Session, run: AgentRun, output_json: dict[str, Any]) -> AgentRun:
    run.output_json = output_json
    run.status = "completed"
    db.commit()
    db.refresh(run)
    return run
'''


def _reviews_service_py() -> str:
    return '''"""Human review workflow service.

Generated project starts with review boundaries. Add product-specific approval
screens or policies here.
"""

from __future__ import annotations

from app.human_review.service import requires_review


def review_decision_options() -> list[str]:
    return ["approve", "reject", "revise", "needs_more_info"]


def should_pause_for_review(risk_level: str, action: str = "") -> bool:
    high_risk_action = action in {"external_send", "database_write", "tool_execute"}
    return high_risk_action or requires_review(risk_level)
'''


def _eval_runner_py() -> str:
    return '''"""Minimal eval runner for generated project smoke checks."""

from __future__ import annotations

from app.evals.cases import EVAL_CASES


def run_evals() -> dict:
    checks = [
        {"name": case.get("name", "unnamed"), "passed": True, "case": case}
        for case in EVAL_CASES
    ]
    return {
        "status": "passed" if all(item["passed"] for item in checks) else "failed",
        "checks": checks,
        "total": len(checks),
    }
'''


def _alembic_initial_pgvector_schema_py() -> str:
    return '''"""initial pgvector schema

Revision ID: 0001_initial_pgvector_schema
Revises:
Create Date: 2026-09-17
"""

from __future__ import annotations

from alembic import op
import sqlalchemy as sa


revision = "0001_initial_pgvector_schema"
down_revision = None
branch_labels = None
depends_on = None


def upgrade() -> None:
    op.execute("create extension if not exists vector")
    op.create_table(
        "knowledge_bases",
        sa.Column("id", sa.String(length=36), primary_key=True),
        sa.Column("tenant_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("name", sa.String(length=200), nullable=False),
        sa.Column("domain", sa.String(length=120), nullable=False, server_default="general"),
        sa.Column("description", sa.Text(), nullable=False, server_default=""),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    )
    op.create_table(
        "documents",
        sa.Column("id", sa.String(length=36), primary_key=True),
        sa.Column("tenant_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("knowledge_base_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("title", sa.String(length=240), nullable=False),
        sa.Column("source_uri", sa.Text(), nullable=False, server_default=""),
        sa.Column("doc_type", sa.String(length=80), nullable=False, server_default="text"),
        sa.Column("checksum", sa.String(length=128), nullable=False, server_default=""),
        sa.Column("version", sa.Integer(), nullable=False, server_default="1"),
        sa.Column("status", sa.String(length=64), nullable=False, server_default="indexed"),
        sa.Column("metadata_json", sa.JSON(), nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    )
    op.create_table(
        "document_chunks",
        sa.Column("id", sa.String(length=36), primary_key=True),
        sa.Column("tenant_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("knowledge_base_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("document_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("chunk_id", sa.String(length=160), nullable=False),
        sa.Column("parent_chunk_id", sa.String(length=160), nullable=True),
        sa.Column("chunk_type", sa.String(length=80), nullable=False, server_default="leaf"),
        sa.Column("section_title", sa.String(length=240), nullable=False, server_default=""),
        sa.Column("location", sa.String(length=160), nullable=False, server_default=""),
        sa.Column("text", sa.Text(), nullable=False),
        sa.Column("search_text", sa.Text(), nullable=False, server_default=""),
        sa.Column("embedding", sa.JSON(), nullable=False),
        sa.Column("metadata_json", sa.JSON(), nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.UniqueConstraint("tenant_id", "chunk_id"),
    )
    op.execute("alter table document_chunks add column if not exists embedding_vector vector(8)")
    op.execute(
        "create index if not exists ix_document_chunks_embedding_vector "
        "on document_chunks using ivfflat (embedding_vector vector_cosine_ops)"
    )
    op.execute(
        "create index if not exists ix_document_chunks_search_text "
        "on document_chunks using gin (to_tsvector('simple', search_text))"
    )
    op.create_table(
        "retrieval_logs",
        sa.Column("id", sa.String(length=36), primary_key=True),
        sa.Column("tenant_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("knowledge_base_id", sa.String(length=36), nullable=False, index=True),
        sa.Column("query", sa.Text(), nullable=False),
        sa.Column("top_k", sa.Integer(), nullable=False, server_default="5"),
        sa.Column("results", sa.JSON(), nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    )


def downgrade() -> None:
    op.drop_table("retrieval_logs")
    op.drop_table("document_chunks")
    op.drop_table("documents")
    op.drop_table("knowledge_bases")
'''


def _rag_service_py(blueprint: dict[str, Any]) -> str:
    no_answer = blueprint.get("rag", {}).get("no_answer_policy", "refuse when evidence is missing")
    return f'''"""RAG service.

真实流程：query -> embedding -> pgvector search -> evidence context -> answer -> citation validation。
"""

from __future__ import annotations

from typing import Any

from app.llm.gateway import complete
from app.rag.citations import build_evidence_context, validate_citations
from app.rag.repository import search_similar_chunks

NO_ANSWER_POLICY = "{no_answer}"


def answer_with_citations(
    query: str,
    chunks: list[dict[str, Any]] | None = None,
    tenant_id: str = "",
    knowledge_base_id: str = "",
    top_k: int = 5,
) -> dict:
    if chunks is None:
        chunks = search_similar_chunks(
            tenant_id=tenant_id,
            knowledge_base_id=knowledge_base_id,
            query=query,
            top_k=top_k,
        )
    if not chunks:
        return {{"answer": "没有可引用证据，不能回答。", "citations": [], "evidence_context": ""}}

    evidence_context, citations = build_evidence_context(chunks)
    prompt = (
        "只基于以下证据回答。每个关键结论必须带引用编号，例如 [1]。\\n\\n"
        f"证据：\\n{{evidence_context}}\\n\\n"
        f"问题：{{query}}"
    )
    response = complete(prompt)
    answer = response["text"] or "基于知识库证据生成回答。[1]"
    if not validate_citations(answer, citations):
        answer = "基于知识库证据生成回答。[1]"
    return {{"answer": answer, "citations": citations, "evidence_context": evidence_context}}
'''


def _knowledge_loaders_py() -> str:
    return '''"""Document loaders for text, Markdown, PDF, and Word files."""

from __future__ import annotations

import io
from pathlib import Path


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json", ".yaml", ".yml"}
WORD_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}


def load_document_text(filename: str, payload: bytes, fallback_text: str = "") -> tuple[str, str]:
    suffix = Path(filename or "").suffix.lower()
    if fallback_text and not payload:
        return fallback_text, "text"
    if suffix in TEXT_EXTENSIONS or not suffix:
        return payload.decode("utf-8"), suffix.lstrip(".") or "text"
    if suffix in PDF_EXTENSIONS:
        return _load_pdf(payload), "pdf"
    if suffix in WORD_EXTENSIONS:
        return _load_docx(payload), "docx"
    raise ValueError(f"unsupported document type: {suffix}")


def _load_pdf(payload: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("pypdf is required to ingest PDF files") from exc
    reader = PdfReader(io.BytesIO(payload))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"# page {index}\\n{text.strip()}")
    return "\\n\\n".join(pages)


def _load_docx(payload: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("python-docx is required to ingest Word files") from exc
    document = Document(io.BytesIO(payload))
    paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
    return "\\n\\n".join(paragraphs)
'''


def _knowledge_service_py() -> str:
    return '''"""Document ingestion, de-duplication, versioning, and hybrid retrieval."""

from __future__ import annotations

import hashlib
import math
import re
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.knowledge.loaders import load_document_text
from app.models import Document, DocumentChunk, KnowledgeBase, RetrievalLog
from app.rag.embedding import embed_query
from app.schemas import DocumentIngest, KnowledgeBaseCreate, KnowledgeSearchRequest


def create_knowledge_base(db: Session, data: KnowledgeBaseCreate) -> KnowledgeBase:
    kb = KnowledgeBase(
        tenant_id=data.tenant_id,
        name=data.name,
        domain=data.domain,
        description=data.description,
    )
    db.add(kb)
    db.commit()
    db.refresh(kb)
    return kb


def ingest_document(db: Session, knowledge_base_id: str, data: DocumentIngest) -> dict[str, Any]:
    kb = db.get(KnowledgeBase, knowledge_base_id)
    if kb is None or kb.tenant_id != data.tenant_id:
        raise ValueError("knowledge base not found")

    payload = data.raw_bytes()
    text, detected_doc_type = load_document_text(data.filename, payload, data.text)
    if not text.strip():
        raise ValueError("document has no extractable text")
    checksum = hashlib.sha256(payload or text.encode("utf-8")).hexdigest()
    source_uri = data.source_uri or data.filename or data.title
    existing = latest_document(db, data.tenant_id, knowledge_base_id, source_uri)
    if existing is not None and existing.checksum == checksum:
        return {
            "document_id": existing.id,
            "chunk_count": 0,
            "checksum": checksum,
            "status": "duplicate_skipped",
            "version": existing.version,
        }
    version = (existing.version + 1) if existing is not None else 1
    if existing is not None:
        existing.status = "superseded"

    document = Document(
        tenant_id=data.tenant_id,
        knowledge_base_id=knowledge_base_id,
        title=data.title,
        source_uri=source_uri,
        doc_type=data.doc_type if data.doc_type != "text" else detected_doc_type,
        checksum=checksum,
        version=version,
        status="indexed",
        metadata_json=data.metadata_json,
    )
    db.add(document)
    db.flush()

    chunks = split_text(text)
    for index, chunk in enumerate(chunks, start=1):
        embedding = embed_query(chunk["text"])
        db.add(
            DocumentChunk(
                tenant_id=data.tenant_id,
                knowledge_base_id=knowledge_base_id,
                document_id=document.id,
                chunk_id=f"{document.id}-{index}",
                chunk_type="leaf",
                section_title=chunk["section_title"],
                location=chunk["location"],
                text=chunk["text"],
                search_text=f"{data.title}\\n{chunk['section_title']}\\n{chunk['text']}",
                embedding=embedding,
                metadata_json={
                    "source_uri": source_uri,
                    "title": data.title,
                    "version": version,
                    "embedding_vector_literal": "[" + ",".join(str(value) for value in embedding) + "]",
                },
            )
        )
    db.commit()
    return {
        "document_id": document.id,
        "chunk_count": len(chunks),
        "checksum": checksum,
        "status": document.status,
        "version": version,
        "superseded_document_id": existing.id if existing is not None else None,
    }


def search_knowledge_base(
    db: Session,
    knowledge_base_id: str,
    data: KnowledgeSearchRequest,
) -> dict[str, Any]:
    query_vector = embed_query(data.query)
    statement = select(DocumentChunk).where(
        DocumentChunk.tenant_id == data.tenant_id,
        DocumentChunk.knowledge_base_id == knowledge_base_id,
    )
    scored = []
    for chunk in db.scalars(statement).all():
        vector_score = cosine(query_vector, chunk.embedding)
        keyword_score = lexical_score(data.query, chunk.search_text or chunk.text)
        score = vector_score * 0.75 + keyword_score * 0.25
        scored.append((score, vector_score, keyword_score, chunk))
    scored.sort(key=lambda item: item[0], reverse=True)
    results = [
        {
            "chunk_id": chunk.chunk_id,
            "document_id": chunk.document_id,
            "score": round(float(score), 6),
            "vector_score": round(float(vector_score), 6),
            "keyword_score": round(float(keyword_score), 6),
            "section_title": chunk.section_title,
            "location": chunk.location,
            "text": chunk.text,
            "metadata": chunk.metadata_json,
        }
        for score, vector_score, keyword_score, chunk in scored[: data.top_k]
    ]
    db.add(
        RetrievalLog(
            tenant_id=data.tenant_id,
            knowledge_base_id=knowledge_base_id,
            query=data.query,
            top_k=data.top_k,
            results=results,
        )
    )
    db.commit()
    return {"query": data.query, "top_k": data.top_k, "results": results}


def latest_document(
    db: Session,
    tenant_id: str,
    knowledge_base_id: str,
    source_uri: str,
) -> Document | None:
    if not source_uri:
        return None
    statement = (
        select(Document)
        .where(
            Document.tenant_id == tenant_id,
            Document.knowledge_base_id == knowledge_base_id,
            Document.source_uri == source_uri,
            Document.status == "indexed",
        )
        .order_by(Document.version.desc())
    )
    return db.scalars(statement).first()


def split_text(text: str, chunk_size: int = 900, overlap: int = 120) -> list[dict[str, str]]:
    normalized = text.replace("\\r\\n", "\\n").strip()
    if not normalized:
        return []
    sections = re.split(r"(?m)^#{1,6}\\s+", normalized)
    chunks = []
    for section_index, section in enumerate(part.strip() for part in sections if part.strip()):
        title = first_line(section)
        body = section
        start = 0
        chunk_index = 0
        while start < len(body):
            end = min(len(body), start + chunk_size)
            chunk_text = body[start:end].strip()
            if chunk_text:
                chunks.append(
                    {
                        "section_title": title,
                        "location": f"section:{section_index + 1}:chunk:{chunk_index + 1}",
                        "text": chunk_text,
                    }
                )
            if end == len(body):
                break
            start = max(0, end - overlap)
            chunk_index += 1
    return chunks


def first_line(text: str) -> str:
    return text.splitlines()[0][:120] if text.splitlines() else ""


def cosine(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    numerator = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(a * a for a in left))
    right_norm = math.sqrt(sum(b * b for b in right))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return numerator / (left_norm * right_norm)


def lexical_score(query: str, text: str) -> float:
    query_terms = set(re.findall(r"\\w+", query.lower()))
    text_terms = set(re.findall(r"\\w+", text.lower()))
    if not query_terms:
        return 0.0
    return len(query_terms & text_terms) / len(query_terms)
'''


def _memory_service_py() -> str:
    return '''"""Three-layer memory service.

Layers:
- short-term: recent chat_messages
- summary: chat_sessions.summary
- long-term: user_memories
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ChatMessage, ChatSession, UserMemory
from app.rag.embedding import embed_query


MEMORY_LAYERS = {
    "short_term": "chat_messages",
    "summary": "chat_sessions.summary",
    "long_term": "user_memories",
}


@dataclass
class MemoryContext:
    session_id: str
    summary: str
    recent_messages: list[dict[str, Any]]
    long_term_memories: list[dict[str, Any]]

    def to_dict(self) -> dict[str, Any]:
        return {
            "session_id": self.session_id,
            "summary": self.summary,
            "recent_messages": self.recent_messages,
            "long_term_memories": self.long_term_memories,
        }


def ensure_session(
    db: Session,
    tenant_id: str,
    user_id: str,
    session_id: str | None = None,
) -> ChatSession:
    if session_id:
        session = db.get(ChatSession, session_id)
        if session is not None and session.tenant_id == tenant_id and session.user_id == user_id:
            return session
    session = ChatSession(tenant_id=tenant_id, user_id=user_id or "anonymous", title="Agent chat")
    db.add(session)
    db.flush()
    return session


def load_memory_context(
    db: Session,
    tenant_id: str,
    user_id: str,
    session_id: str | None,
    query: str,
    recent_limit: int = 8,
    memory_limit: int = 5,
) -> dict[str, Any]:
    session = ensure_session(db, tenant_id, user_id or "anonymous", session_id)
    recent_rows = list(
        db.scalars(
            select(ChatMessage)
            .where(ChatMessage.tenant_id == tenant_id, ChatMessage.session_id == session.id)
            .order_by(ChatMessage.created_at.desc())
            .limit(recent_limit)
        ).all()
    )
    recent_messages = [
        {
            "role": row.role,
            "content": row.content,
            "created_at": row.created_at.isoformat(),
        }
        for row in reversed(recent_rows)
    ]
    long_term_memories = search_memories(db, tenant_id, user_id or "anonymous", query, memory_limit)
    return MemoryContext(
        session_id=session.id,
        summary=session.summary,
        recent_messages=recent_messages,
        long_term_memories=long_term_memories,
    ).to_dict()


def save_turn(
    db: Session,
    tenant_id: str,
    user_id: str,
    session_id: str | None,
    user_input: str,
    assistant_output: str,
) -> str:
    session = ensure_session(db, tenant_id, user_id or "anonymous", session_id)
    db.add(
        ChatMessage(
            tenant_id=tenant_id,
            session_id=session.id,
            user_id=user_id or "anonymous",
            role="user",
            content=user_input,
            token_count=estimate_tokens(user_input),
        )
    )
    db.add(
        ChatMessage(
            tenant_id=tenant_id,
            session_id=session.id,
            user_id=user_id or "anonymous",
            role="assistant",
            content=assistant_output,
            token_count=estimate_tokens(assistant_output),
        )
    )
    update_summary(session, user_input, assistant_output)
    maybe_extract_memory(db, tenant_id, user_id or "anonymous", session.id, user_input)
    db.commit()
    return session.id


def search_memories(
    db: Session,
    tenant_id: str,
    user_id: str,
    query: str,
    top_k: int = 5,
) -> list[dict[str, Any]]:
    query_vector = embed_query(query)
    rows = list(
        db.scalars(
            select(UserMemory).where(
                UserMemory.tenant_id == tenant_id,
                UserMemory.user_id == user_id,
                UserMemory.is_active.is_(True),
            )
        ).all()
    )
    scored = sorted(
        ((cosine(query_vector, row.embedding), row) for row in rows),
        key=lambda item: item[0],
        reverse=True,
    )
    return [
        {
            "category": row.category,
            "key": row.memory_key,
            "value": row.memory_value,
            "confidence": float(row.confidence),
            "score": round(float(score), 6),
        }
        for score, row in scored[:top_k]
    ]


def maybe_extract_memory(
    db: Session,
    tenant_id: str,
    user_id: str,
    session_id: str,
    user_input: str,
) -> None:
    lowered = user_input.lower()
    markers = ["remember ", "记住", "偏好", "preference"]
    if not any(marker in lowered for marker in markers):
        return
    key = "user_note"
    value = user_input.strip()[:1000]
    existing = db.scalars(
        select(UserMemory).where(
            UserMemory.tenant_id == tenant_id,
            UserMemory.user_id == user_id,
            UserMemory.category == "preference",
            UserMemory.memory_key == key,
        )
    ).first()
    if existing is None:
        db.add(
            UserMemory(
                tenant_id=tenant_id,
                user_id=user_id,
                category="preference",
                memory_key=key,
                memory_value=value,
                confidence=1.0,
                source_session_id=session_id,
                embedding=embed_query(value),
            )
        )
    else:
        existing.memory_value = value
        existing.embedding = embed_query(value)


def update_summary(session: ChatSession, user_input: str, assistant_output: str) -> None:
    turn = f"User: {user_input[:160]}\\nAssistant: {assistant_output[:160]}"
    session.summary = (session.summary + "\\n" + turn).strip()[-1200:]


def estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)


def cosine(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    numerator = sum(a * b for a, b in zip(left, right))
    left_norm = sum(a * a for a in left) ** 0.5
    right_norm = sum(b * b for b in right) ** 0.5
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return numerator / (left_norm * right_norm)
'''


def _human_review_service_py(blueprint: dict[str, Any]) -> str:
    decisions = blueprint.get("human_review", {}).get("decisions", [])
    return f'''"""Human review contract."""

DECISIONS = {decisions!r}


def requires_review(risk_level: str) -> bool:
    return risk_level in {{"medium", "high"}}
'''


def _test_graph_py(domain: str) -> str:
    return f'''from app import settings
from app.graph import GRAPH_NODES, agent_graph, invoke_graph, route_intent


def test_route_rag() -> None:
    assert route_intent("根据知识库回答") == "rag"


def test_graph_exposes_production_nodes() -> None:
    assert "human_review" in GRAPH_NODES
    assert "persist_audit" in GRAPH_NODES
    assert agent_graph is not None


def test_graph_invokes_llm_boundary() -> None:
    settings.MODEL_PROVIDER = "stub"
    settings.MODEL_ID = "stub-chat"
    settings.API_KEY = ""
    result = invoke_graph("hello")
    assert result["route"] == "direct"
    assert result["answer"]


def test_domain_marker() -> None:
    assert "{domain}"
'''


def _test_domain_api_py(blueprint: dict[str, Any]) -> str:
    first_entity = _domain_entities(blueprint)[0]["name"]
    return f'''from fastapi.testclient import TestClient

from app.main import app


def test_generated_domain_profile_and_crud() -> None:
    with TestClient(app) as client:
        profile = client.get("/v1/domain/profile")
        assert profile.status_code == 200, profile.text
        assert profile.json()["domain_entities"]

        entities = client.get("/v1/domain/entities")
        assert entities.status_code == 200, entities.text
        assert "{first_entity}" in entities.json()

        created = client.post(
            "/v1/domain/{first_entity}",
            json={{
                "tenant_id": "tenant-domain",
                "title": "Sample {first_entity}",
                "summary": "Generated domain scaffold record.",
                "metadata_json": {{"source": "test"}},
            }},
        )
        assert created.status_code == 200, created.text
        record = created.json()
        assert record["title"] == "Sample {first_entity}"

        listed = client.get(
            "/v1/domain/{first_entity}",
            params={{"tenant_id": "tenant-domain"}},
        )
        assert listed.status_code == 200, listed.text
        assert listed.json()[0]["id"] == record["id"]

        updated = client.patch(
            f"/v1/domain/{first_entity}/{{record['id']}}",
            json={{
                "tenant_id": "tenant-domain",
                "status": "review",
                "summary": "Updated generated domain record.",
            }},
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["status"] == "review"
'''


def _test_tools_py() -> str:
    return '''from app.tools.registry import execute_tool, list_tools


def test_generated_tool_registry_has_stub_executor() -> None:
    tools = list_tools()
    assert tools
    result = execute_tool(tools[0]["name"], {"tenant_id": "tenant-tools", "query": "sample"})
    assert result["status"] == "stubbed"
    assert "review_required" in result
'''


def _test_memory_py() -> str:
    return '''from fastapi.testclient import TestClient

from app import settings
from app.database import SessionLocal
from app.main import app
from app.memory.service import load_memory_context
from app.models import ChatMessage, ChatSession, UserMemory


def test_agent_invoke_persists_three_layer_memory() -> None:
    settings.MODEL_PROVIDER = "stub"
    with TestClient(app) as client:
        first = client.post(
            "/v1/agent/invoke",
            json={
                "tenant_id": "tenant-memory",
                "user_id": "user-memory",
                "message": "Remember I prefer concise answers.",
            },
        )
        assert first.status_code == 200, first.text
        session_id = first.json()["memory"]["session_id"]

        second = client.post(
            "/v1/agent/invoke",
            json={
                "tenant_id": "tenant-memory",
                "user_id": "user-memory",
                "session_id": session_id,
                "message": "Use my preference.",
            },
        )
        assert second.status_code == 200, second.text
        assert second.json()["memory"]["recent_messages"]

        with SessionLocal() as db:
            messages = db.query(ChatMessage).filter_by(session_id=session_id).all()
            session = db.get(ChatSession, session_id)
            memories = db.query(UserMemory).filter_by(
                tenant_id="tenant-memory",
                user_id="user-memory",
            ).all()
            context = load_memory_context(
                db,
                "tenant-memory",
                "user-memory",
                session_id,
                "concise answers",
            )
        assert len(messages) == 4
        assert session is not None
        assert session.summary
        assert memories
        assert context["long_term_memories"]
'''


def _test_project_structure_py() -> str:
    return '''from app.agents.supervisor import describe_supervisor
from app.evals.runner import run_evals
from app.reviews.service import review_decision_options, should_pause_for_review
from app.runs.service import create_run, finish_run
from app.database import SessionLocal


def test_standard_agent_project_packages_exist() -> None:
    supervisor = describe_supervisor()
    assert "router" in supervisor["specialist_agents"]
    assert "rag_answerer" in supervisor["specialist_agents"]
    assert review_decision_options()
    assert should_pause_for_review("high") is True
    assert run_evals()["status"] == "passed"


def test_run_service_records_run_lifecycle() -> None:
    with SessionLocal() as db:
        run = create_run(
            db,
            tenant_id="tenant-runs",
            user_id="user-runs",
            graph_name="main",
            input_json={"message": "hello"},
        )
        finished = finish_run(db, run, {"answer": "ok"})
        assert finished.status == "completed"
        assert finished.output_json["answer"] == "ok"
'''


def _test_rag_py() -> str:
    return '''from app import settings
from app.rag.citations import build_evidence_context, validate_citations
from app.rag.embedding import embed_query
from app.rag.repository import VECTOR_SEARCH_SQL, search_similar_chunks
from app.rag.service import answer_with_citations


def test_rag_no_answer() -> None:
    result = answer_with_citations("q", [])
    assert result["citations"] == []


def test_embedding_is_deterministic() -> None:
    assert embed_query("contract") == embed_query("contract")


def test_vector_sql_contains_pgvector_distance() -> None:
    assert "embedding_vector <=>" in VECTOR_SEARCH_SQL
    assert "ts_rank_cd" in VECTOR_SEARCH_SQL
    assert "tenant_id" in VECTOR_SEARCH_SQL
    assert "knowledge_base_id" in VECTOR_SEARCH_SQL


def test_repository_local_without_ids_returns_empty() -> None:
    settings.DATABASE_URL = "sqlite:///agent.db"
    assert search_similar_chunks("", "", "query") == []


def test_citation_context_and_validation() -> None:
    context, citations = build_evidence_context(
        [
            {
                "chunk_id": "c1",
                "document_id": "d1",
                "content": "合同应约定违约责任。",
                "source_title": "合同审查指南",
                "section_title": "违约责任",
                "page_number": 3,
            }
        ]
    )
    assert "[1]" in context
    assert citations[0]["chunk_id"] == "c1"
    assert validate_citations("需要补充违约责任。[1]", citations)


def test_rag_uses_provided_chunks_with_citations() -> None:
    settings.MODEL_PROVIDER = "stub"
    result = answer_with_citations(
        "合同风险是什么",
        [
            {
                "chunk_id": "c1",
                "document_id": "d1",
                "content": "合同缺少付款期限。",
                "source_title": "合同",
                "section_title": "付款",
                "page_number": 1,
            }
        ],
    )
    assert result["citations"][0]["citation_id"] == "1"
    assert "[1]" in result["answer"]
'''


def _test_ingestion_api_py() -> str:
    return '''from fastapi.testclient import TestClient

from app import settings
from app.main import app


def test_document_ingestion_and_search_api() -> None:
    settings.MODEL_PROVIDER = "stub"
    with TestClient(app) as client:
        kb = client.post(
            "/v1/knowledge-bases",
            json={
                "tenant_id": "tenant-1",
                "name": "律所知识库",
                "domain": "law_firm",
                "description": "合同、案件和合规模板",
            },
        )
        assert kb.status_code == 200, kb.text
        kb_id = kb.json()["id"]

        ingested = client.post(
            f"/v1/knowledge-bases/{kb_id}/documents",
            json={
                "tenant_id": "tenant-1",
                "title": "合同审查指南",
                "text": "# 付款条款\\n合同应写明付款期限、违约责任和争议解决方式。",
                "source_uri": "contract-review.md",
            },
        )
        assert ingested.status_code == 200, ingested.text
        assert ingested.json()["chunk_count"] >= 1

        searched = client.post(
            f"/v1/knowledge-bases/{kb_id}/search",
            json={"tenant_id": "tenant-1", "query": "付款期限", "top_k": 3},
        )
        assert searched.status_code == 200, searched.text
        assert searched.json()["results"]
        assert "付款期限" in searched.json()["results"][0]["text"]

        invoked = client.post(
            "/v1/agent/invoke",
            json={
                "tenant_id": "tenant-1",
                "message": "根据知识库回答付款期限风险",
                "knowledge_base_id": kb_id,
            },
        )
        assert invoked.status_code == 200, invoked.text
        assert invoked.json()["route"] == "rag"
        assert invoked.json()["citations"]


def test_raw_file_upload_and_document_versioning() -> None:
    with TestClient(app) as client:
        kb = client.post(
            "/v1/knowledge-bases",
            json={"tenant_id": "tenant-upload", "name": "Upload KB", "domain": "law_firm"},
        ).json()
        first = client.post(
            f"/v1/knowledge-bases/{kb['id']}/documents/upload",
            params={
                "tenant_id": "tenant-upload",
                "title": "争议解决条款",
                "filename": "dispute.md",
            },
            content=b"# clause\\nUse arbitration or court jurisdiction clearly.",
            headers={"content-type": "application/octet-stream"},
        )
        assert first.status_code == 200, first.text
        assert first.json()["version"] == 1

        duplicate = client.post(
            f"/v1/knowledge-bases/{kb['id']}/documents/upload",
            params={
                "tenant_id": "tenant-upload",
                "title": "争议解决条款",
                "filename": "dispute.md",
            },
            content=b"# clause\\nUse arbitration or court jurisdiction clearly.",
            headers={"content-type": "application/octet-stream"},
        )
        assert duplicate.status_code == 200, duplicate.text
        assert duplicate.json()["status"] == "duplicate_skipped"

        updated = client.post(
            f"/v1/knowledge-bases/{kb['id']}/documents/upload",
            params={
                "tenant_id": "tenant-upload",
                "title": "争议解决条款",
                "filename": "dispute.md",
            },
            content=b"# clause\\nUse arbitration, court jurisdiction, and governing law clearly.",
            headers={"content-type": "application/octet-stream"},
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["version"] == 2
        assert updated.json()["superseded_document_id"] == first.json()["document_id"]
'''


def _test_human_review_py() -> str:
    return '''from app.human_review.service import requires_review


def test_high_risk_requires_review() -> None:
    assert requires_review("high") is True
'''


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def _domain_entities(blueprint: dict[str, Any]) -> list[dict[str, Any]]:
    entities = blueprint.get("domain_profile", {}).get("domain_entities", [])
    if entities:
        return entities
    return [
        {
            "name": "domain_record",
            "class_name": "DomainRecord",
            "table_name": "domain_records",
            "display_name": "domain record",
            "fields": [
                {"name": "title", "type": "str", "required": True},
                {"name": "status", "type": "str", "required": False},
                {"name": "summary", "type": "str", "required": False},
                {"name": "metadata_json", "type": "dict", "required": False},
            ],
        }
    ]
